import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

OUT_PATH = "orch_metadata.xlsx"

# (sheet_name, tab_color_hex) in the order they should appear (parents before children)
SHEETS = [
    ("ObjectIDs", "BF8F00"),
    ("TaskType",  "7030A0"),
    ("Jobs",      "2F5597"),
    ("Tasks",     "548235"),
]

# column spec: (name, sql_type, nullable, default, key_note)
COLUMNS = {
    "ObjectIDs": [
        ("WorkspaceName", "VARCHAR(200)", False, None, "PK (1/2)"),
        ("ObjectName",    "VARCHAR(200)", False, None, "PK (2/2)"),
        ("ObjectID",      "VARCHAR(200)", False, None, None),
        ("WorkspaceID",   "VARCHAR(200)", True,  None, None),
    ],
    "TaskType": [
        ("TaskType",      "NVARCHAR(50)", False, None, "PK"),
        ("AllowParallel", "BIT",          False, "1",  None),
    ],
    "Jobs": [
        ("JobName",                "VARCHAR(200)",  False, None, "PK"),
        ("Include",                "BIT",           False, "1",  None),
        ("TimeoutInSeconds",       "INT",           False, None, None),
        ("Retries",                "INT",           False, None, None),
        ("RetryIntervalInSeconds", "INT",           False, None, None),
        ("ScheduledStartUTC",      "TIME(0)",       True,  None, None),
        ("ParametersJson",         "NVARCHAR(MAX)", True,  None, None),
        ("Dependencies",           "NVARCHAR(MAX)", True,  None, None),
        ("WorkspaceName",          "VARCHAR(200)",  True,  None, "FK -> orch.ObjectIDs (WorkspaceName,JobName)=(WorkspaceName,ObjectName)"),
        ("Environment",            "NVARCHAR(100)", True,  None, None),
        ("LoggingLevel",           "TINYINT",       False, "1",  "FK -> log.LoggingLevel"),
    ],
    "Tasks": [
        ("TaskName",               "NVARCHAR(200)", False, None, "PK"),
        ("Include",                "BIT",           False, "1",  None),
        ("JobName",                "VARCHAR(200)",  False, None, "FK -> orch.Jobs"),
        ("ObjectName",             "VARCHAR(200)",  False, None, "FK -> orch.ObjectIDs (composite w/ WorkspaceName)"),
        ("WorkspaceName",          "VARCHAR(200)",  True,  None, "FK -> orch.ObjectIDs (composite)"),
        ("TimeoutInSeconds",       "INT",           False, None, None),
        ("Retries",                "INT",           False, None, None),
        ("RetryIntervalInSeconds", "INT",           False, None, None),
        ("ParametersJson",         "NVARCHAR(MAX)", True,  None, None),
        ("Dependencies",           "NVARCHAR(MAX)", True,  None, None),
        ("TaskType",               "NVARCHAR(50)",  False, None, "FK -> orch.TaskType"),
        ("System",                 "NVARCHAR(100)", True,  None, None),
        ("Layer",                  "NVARCHAR(100)", True,  None, None),
        ("LoggingLevel",           "TINYINT",       False, "1",  "FK -> log.LoggingLevel"),
    ],
}

NAMED_RANGE_ROWS = 1000  # generous headroom for dropdown source ranges

wb = openpyxl.Workbook()
wb.remove(wb.active)

for sheet_name, tab_hex in SHEETS:
    ws = wb.create_sheet(sheet_name)
    ws.sheet_properties.tabColor = tab_hex
    cols = COLUMNS[sheet_name]

    header_fill = PatternFill(start_color=tab_hex, end_color=tab_hex, fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for idx, (name, sql_type, nullable, default, key_note) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=idx, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        note_lines = [f"Type: {sql_type}", f"Nullable: {'Yes' if nullable else 'No'}"]
        if default:
            note_lines.append(f"Default: {default}")
        if key_note:
            note_lines.append(key_note)
        cell.comment = Comment("\n".join(note_lines), "Schema")

        col_letter = get_column_letter(idx)
        width = max(14, len(name) + 4)
        ws.column_dimensions[col_letter].width = min(width, 32)

    # Table region: header + 1 blank body row (Excel auto-extends on next-row entry)
    last_col_letter = get_column_letter(len(cols))
    table_ref = f"A1:{last_col_letter}2"
    table = Table(displayName=f"tbl{sheet_name}", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

# --- Named ranges used as dropdown sources (defined after all sheets exist) ---
wb.defined_names["JobNameList"] = DefinedName("JobNameList", attr_text=f"Jobs!$A$2:$A${NAMED_RANGE_ROWS}")
wb.defined_names["TaskTypeList"] = DefinedName("TaskTypeList", attr_text=f"TaskType!$A$2:$A${NAMED_RANGE_ROWS}")
wb.defined_names["WorkspaceNameList"] = DefinedName("WorkspaceNameList", attr_text=f"ObjectIDs!$A$2:$A${NAMED_RANGE_ROWS}")
wb.defined_names["ObjectNameList"] = DefinedName("ObjectNameList", attr_text=f"ObjectIDs!$B$2:$B${NAMED_RANGE_ROWS}")

BOOL_DV_TARGETS = {
    "Jobs": ["Include"],
    "Tasks": ["Include"],
    "TaskType": ["AllowParallel"],
}
FK_DV_TARGETS = {
    "Jobs": [("WorkspaceName", "WorkspaceNameList")],
    "Tasks": [("JobName", "JobNameList"), ("TaskType", "TaskTypeList"),
              ("WorkspaceName", "WorkspaceNameList"), ("ObjectName", "ObjectNameList")],
}

for sheet_name, tab_hex in SHEETS:
    ws = wb[sheet_name]
    cols = COLUMNS[sheet_name]
    col_index = {name: i + 1 for i, (name, *_rest) in enumerate(cols)}

    for bool_col in BOOL_DV_TARGETS.get(sheet_name, []):
        idx = col_index[bool_col]
        letter = get_column_letter(idx)
        dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True, showDropDown=False)
        dv.add(f"{letter}2:{letter}{NAMED_RANGE_ROWS}")
        ws.add_data_validation(dv)

    for fk_col, range_name in FK_DV_TARGETS.get(sheet_name, []):
        idx = col_index[fk_col]
        letter = get_column_letter(idx)
        dv = DataValidation(type="list", formula1=range_name, allow_blank=True, showDropDown=False)
        dv.add(f"{letter}2:{letter}{NAMED_RANGE_ROWS}")
        ws.add_data_validation(dv)

wb.save(OUT_PATH)
print("Workbook written:", OUT_PATH)
